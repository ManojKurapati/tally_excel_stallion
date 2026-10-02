"""Excel workbook for a Tally expert to verify extracted data against TallyPrime.

One workbook per company. Every figure the expert can check has the TallyPrime
screen where it is found, a blank column for the value seen in Tally and a
formula that flags mismatches. Amounts are shown in separate Debit / Credit
columns, the way Tally shows them.
"""

from __future__ import annotations

import os
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from stallion_tally import __version__
from stallion_tally.reports.checks import (
    ZERO,
    CheckResult,
    counts_in_books,
    ledger_movements,
    period_totals,
    postings_by_voucher,
    run_checks,
    split_dr_cr,
    voucher_totals,
    voucher_type_summaries,
)
from stallion_tally.reports.loader import CompanyReportData

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_TEXT = 32_767

MONEY = "#,##0.00"
QTY = "#,##0.####"
DATE = "d-mmm-yyyy"
INT = "#,##0"

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_EXPERT_HEADER_FILL = PatternFill("solid", fgColor="BF8F00")
_EXPERT_FILL = PatternFill("solid", fgColor="FFF2CC")
_NOTE_FONT = Font(italic=True, color="595959")
_SECTION_FONT = Font(bold=True, size=12, color="1F4E78")
_STATUS_FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "WARN": PatternFill("solid", fgColor="FFEB9C"),
    "INFO": PatternFill("solid", fgColor="DDEBF7"),
}
_WRAP = Alignment(wrap_text=True, vertical="top")

TALLY_PATHS = {
    "groups": "Gateway of Tally > Chart of Accounts > Groups",
    "ledgers": "Gateway of Tally > Chart of Accounts > Ledgers",
    "stock_items": "Gateway of Tally > Chart of Accounts > Stock Items",
    "opening": "Gateway of Tally > Trial Balance (F12 Configure: show Opening Balance), "
    "period starting on 'Books from' date",
    "day_book": "Gateway of Tally > Day Book (Alt+F2: set the period above)",
    "statistics": "Gateway of Tally > Display More Reports > Statement of Accounts > "
    "Statistics (Alt+F2: same period)",
    "receivables": "Gateway of Tally > Display More Reports > Statement of Accounts > "
    "Outstandings > Receivables",
    "payables": "Gateway of Tally > Display More Reports > Statement of Accounts > "
    "Outstandings > Payables",
    "ledger_vouchers": "Gateway of Tally > Display More Reports > Account Books > Ledger > "
    "(select ledger), same period",
}


class Formula(str):
    """A value deliberately written as an Excel formula (all other text is literal)."""


@dataclass(frozen=True)
class Col:
    header: str
    width: float = 14
    fmt: str | None = None
    expert: bool = False
    wrap: bool = False


@dataclass
class ExcelReportResult:
    path: Path
    company_name: str
    period_from: date | None
    period_to: date | None
    sheet_rows: dict[str, int] = field(default_factory=dict)
    checks: list[CheckResult] = field(default_factory=list)

    @property
    def failed_checks(self) -> int:
        return sum(1 for c in self.checks if c.status == "FAIL")

    @property
    def warning_checks(self) -> int:
        return sum(1 for c in self.checks if c.status == "WARN")


def tally_date(value: date | None) -> str:
    """Date the way Tally prints it, e.g. 1-Apr-2026."""
    return f"{value.day}-{value:%b-%Y}" if value else "-"


def period_label(data: CompanyReportData) -> str:
    if data.period_from is None and data.period_to is None:
        return "no vouchers"
    return f"{tally_date(data.period_from)} to {tally_date(data.period_to)}"


def safe_file_part(text: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", text).strip(" ._")
    return cleaned[:80] or "company"


def _sheet_title(title: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "", title)[:31]


def _local(value: datetime | None) -> datetime | None:
    # Excel has no time zones; show the operator's local time.
    return value.astimezone().replace(tzinfo=None) if value else None


class SheetWriter:
    """Streams rows into a write-only worksheet, continuing on a new sheet when full."""

    def __init__(
        self,
        workbook: Workbook,
        title: str,
        columns: Sequence[Col],
        note: str | None = None,
        max_rows: int = EXCEL_MAX_ROWS,
    ) -> None:
        self.workbook = workbook
        self.title = title
        self.columns = list(columns)
        self.note = note
        self.max_rows = max_rows
        self.rows_written = 0
        self._part = 0
        self._new_sheet()

    @property
    def next_row(self) -> int:
        """Excel row number the next `append` will write to."""
        return self._row + 1

    def _new_sheet(self) -> None:
        self._part += 1
        title = self.title if self._part == 1 else f"{self.title} ({self._part})"
        self.ws = self.workbook.create_sheet(_sheet_title(title))
        for index, col in enumerate(self.columns, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = col.width
        self._row = 0
        if self.note:
            cell = WriteOnlyCell(self.ws, value=self.note)
            cell.font = _NOTE_FONT
            self.ws.append([cell])
            self._row += 1
        header_row = self._row + 1
        headers = []
        for col in self.columns:
            cell = WriteOnlyCell(self.ws, value=col.header)
            cell.font = _HEADER_FONT
            cell.fill = _EXPERT_HEADER_FILL if col.expert else _HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            headers.append(cell)
        self.ws.append(headers)
        self._row += 1
        self.ws.freeze_panes = f"A{header_row + 1}"
        last = get_column_letter(len(self.columns))
        self.ws.auto_filter.ref = f"A{header_row}:{last}{header_row}"
        self._header_row = header_row

    def _cell(
        self, value: Any, col: Col, fill: PatternFill | None, fmt: str | None
    ) -> WriteOnlyCell:
        literal_text = False
        if isinstance(value, Formula):
            value = str(value)
        elif isinstance(value, bool):
            value = "Yes" if value else "No"
        elif isinstance(value, Decimal):
            value = float(value)
        elif isinstance(value, datetime):
            value = _local(value)
        elif isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_MAX_TEXT]
            literal_text = value.startswith("=")
        cell = WriteOnlyCell(self.ws, value=value)
        if literal_text:
            cell.data_type = "s"  # never let Tally text (e.g. a narration) run as a formula
        fmt = fmt or col.fmt
        if fmt and value is not None:
            cell.number_format = fmt
        if col.wrap:
            cell.alignment = _WRAP
        if fill is not None:
            cell.fill = fill
        elif col.expert:
            cell.fill = _EXPERT_FILL
        return cell

    def append(
        self,
        values: Sequence[Any],
        fills: dict[int, PatternFill] | None = None,
        formats: dict[int, str] | None = None,
    ) -> None:
        """Append one row; `fills`/`formats` override style per column index."""
        if self._row >= self.max_rows:
            self._new_sheet()
        fills = fills or {}
        formats = formats or {}
        padded = list(values) + [None] * (len(self.columns) - len(values))
        self.ws.append(
            [
                self._cell(v, col, fills.get(i), formats.get(i))
                for i, (v, col) in enumerate(zip(padded, self.columns, strict=True))
            ]
        )
        self._row += 1
        self.rows_written += 1

    def extend(self, rows: Iterable[Sequence[Any]]) -> int:
        for row in rows:
            self.append(row)
        return self.rows_written


def _match_formula(extracted_col: str, tally_col: str, row: int) -> Formula:
    return Formula(
        f'=IF({tally_col}{row}="","",'
        f'IF(ABS({extracted_col}{row}-{tally_col}{row})<0.01,"OK","MISMATCH"))'
    )


REMARKS = Col("Expert remarks", 36, expert=True, wrap=True)


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------


def _write_read_me(wb: Workbook, data: CompanyReportData, generated_at: datetime) -> int:
    ws = wb.create_sheet("Read Me")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    company = data.company
    bills_as_on = max((b.snapshot_at for b in data.bills if b.snapshot_at), default=None)

    def section(title: str) -> None:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = _SECTION_FONT
        ws.append([])
        ws.append([cell])

    def line(label: str, value: Any = None) -> None:
        label_cell = WriteOnlyCell(ws, value=label)
        label_cell.font = Font(bold=True)
        if isinstance(value, datetime):
            value = _local(value)
        value_cell = WriteOnlyCell(ws, value=value)
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(value, datetime):
            value_cell.number_format = "d-mmm-yyyy hh:mm"
        ws.append([label_cell, value_cell])

    title = WriteOnlyCell(ws, value=f"Tally data verification — {company.company_name}")
    title.font = Font(bold=True, size=14, color="1F4E78")
    ws.append([title])
    line("Company", company.company_name)
    line("Company GUID in Tally", company.tally_guid or company.company_id)
    line("Books beginning from", tally_date(company.books_from))
    line("Base currency", company.base_currency_name or company.base_currency_symbol or "-")
    line("Voucher period in this workbook", f"{period_label(data)}  ({data.period_source})")
    line(
        "Bills outstanding as on",
        _local(bills_as_on).strftime("%d-%b-%Y %H:%M") if bills_as_on else "-",
    )
    line("Workbook generated at", generated_at)
    line("Generated by", f"stallion-tally {__version__}")

    section("How to verify")
    steps = [
        "1. Open the same company in TallyPrime on the same machine.",
        "2. Go to the 'Summary' sheet. For each row open the Tally screen in the "
        "'Where to find in TallyPrime' column and type the figure you see into the yellow "
        "'Tally figure' column. The 'Match' column turns to OK or MISMATCH by itself.",
        "3. Do the same in the 'Voucher Types' sheet using Tally's Statistics report "
        f"for the period {period_label(data)}.",
        "4. Review the 'Checks' sheet. FAIL = extracted data is internally inconsistent; "
        "WARN = please confirm in Tally. The rows behind each check are in 'Check Details'.",
        "5. Spot-check a sample of vouchers in 'Day Book' / 'Voucher Lines' and a few "
        "ledgers in 'Ledgers' (period Debit/Credit = 'Current Total' in Tally's Ledger "
        "Vouchers report for the same period).",
        "6. Write any finding in the yellow 'Expert remarks' columns and return the file.",
    ]
    for step in steps:
        line("", step)

    section("Conventions")
    for label, text in (
        (
            "Debit / Credit",
            "Amounts are shown as positive numbers in separate Debit and "
            "Credit columns, as in Tally.",
        ),
        (
            "Non-accounting vouchers",
            "Orders, Delivery Notes and Receipt Notes (invoice-view vouchers that are not "
            "invoices) are listed with 'Posts to books' = No and left out of all Debit / Credit "
            "totals and ledger movements, as in Tally.",
        ),
        (
            "Statistics roll-up",
            "Tally's Statistics adds vouchers of types created under a built-in type into that "
            "built-in type's row and the Total (e.g. Journal includes IC JV Expenses). This "
            "workbook counts each voucher under its own type, so compare a built-in type with "
            "its own count plus its sub-types.",
        ),
        (
            "Cancelled vouchers",
            "Listed in the Day Book (they appear in Tally's Day Book) but carry no amounts.",
        ),
        (
            "Optional vouchers",
            "Listed, but excluded from totals and ledger movements because "
            "they do not affect the books.",
        ),
        (
            "Deleted records",
            "Masters and vouchers deleted in Tally since the last sync are not included.",
        ),
        (
            "Bills outstanding",
            "Snapshot taken at the last sync; compare with Tally's Outstandings as on that date.",
        ),
    ):
        line(label, text)

    section("Last successful sync from Tally")
    for dataset in ("groups", "ledgers", "stock_items", "bills", "vouchers"):
        line(dataset.replace("_", " ").title(), data.last_sync.get(dataset) or "never")

    section("Sheets")
    for name, text in (
        ("Summary", "Headline figures to compare with Tally, with the screen to look at."),
        ("Checks", "Automatic consistency checks and their result."),
        ("Check Details", "The individual vouchers / masters behind any FAIL or WARN."),
        ("Voucher Types", "Voucher counts per type, to compare with Tally's Statistics."),
        ("Day Book", "One row per voucher in the period."),
        ("Voucher Lines", "Ledger (accounting) lines of every voucher."),
        ("Inventory Lines", "Stock item lines of every voucher."),
        ("Ledgers", "Ledger master with opening balance and period Debit / Credit."),
        ("Groups", "Group master (Chart of Accounts)."),
        ("Stock Items", "Stock item master with opening stock."),
        ("Bills Outstanding", "Pending bills (receivable and payable)."),
    ):
        line(name, text)
    return 0


def _write_summary(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Summary",
        [
            Col("Area", 18),
            Col("Figure", 34),
            Col("Extracted", 18, MONEY),
            Col("Tally figure", 18, MONEY, expert=True),
            Col("Match", 12),
            Col("Where to find in TallyPrime", 70, wrap=True),
            REMARKS,
        ],
        note=f"Company: {data.company.company_name}   |   Voucher period: {period_label(data)}",
    )
    opening_dr = opening_cr = ZERO
    for ledger in data.ledgers:
        dr, cr = split_dr_cr(ledger.opening_balance)
        opening_dr += dr or ZERO
        opening_cr += cr or ZERO
    period_dr, period_cr = period_totals(data)
    receivable = [b for b in data.bills if b.bill_type == "receivable"]
    payable = [b for b in data.bills if b.bill_type == "payable"]
    period = f" — period {period_label(data)}"

    rows: list[tuple[str, str, Any, str]] = [
        ("Masters", "Number of groups", len(data.groups), TALLY_PATHS["groups"]),
        ("Masters", "Number of ledgers", len(data.ledgers), TALLY_PATHS["ledgers"]),
        ("Masters", "Number of stock items", len(data.stock_items), TALLY_PATHS["stock_items"]),
        ("Opening", "Total opening balance — Debit", opening_dr, TALLY_PATHS["opening"]),
        ("Opening", "Total opening balance — Credit", opening_cr, TALLY_PATHS["opening"]),
        (
            "Day Book",
            "Number of vouchers (excluding optional)",
            sum(not v.is_optional for v in data.vouchers),
            TALLY_PATHS["statistics"] + " — 'Total' row",
        ),
        (
            "Day Book",
            "  of which cancelled",
            sum(v.is_cancelled and not v.is_optional for v in data.vouchers),
            TALLY_PATHS["day_book"] + period,
        ),
        (
            "Day Book",
            "Optional vouchers (not in Statistics)",
            sum(v.is_optional for v in data.vouchers),
            "Gateway of Tally > Display More Reports > Exception Reports > Optional Vouchers",
        ),
        ("Day Book", "Total Debit", period_dr, TALLY_PATHS["day_book"] + period),
        ("Day Book", "Total Credit", period_cr, TALLY_PATHS["day_book"] + period),
        ("Outstanding", "Receivable bills (count)", len(receivable), TALLY_PATHS["receivables"]),
        (
            "Outstanding",
            "Receivable pending amount (Dr)",
            -sum((b.closing_amount or ZERO for b in receivable), ZERO),
            TALLY_PATHS["receivables"],
        ),
        ("Outstanding", "Payable bills (count)", len(payable), TALLY_PATHS["payables"]),
        (
            "Outstanding",
            "Payable pending amount (Cr)",
            sum((b.closing_amount or ZERO for b in payable), ZERO),
            TALLY_PATHS["payables"],
        ),
    ]
    for area, figure, value, where in rows:
        row = sheet.next_row
        fmt = MONEY if isinstance(value, Decimal) else INT
        sheet.append(
            [area, figure, value, None, _match_formula("C", "D", row), where],
            formats={2: fmt, 3: fmt},
        )
    return sheet.rows_written


def _write_checks(wb: Workbook, checks: list[CheckResult]) -> int:
    sheet = SheetWriter(
        wb,
        "Checks",
        [
            Col("Code", 8),
            Col("Check", 48, wrap=True),
            Col("Result", 10),
            Col("Outcome", 40, wrap=True),
            Col("What it means / what to do", 80, wrap=True),
            Col("Expert agrees? (Y/N)", 14, expert=True),
            REMARKS,
        ],
        note="Automatic checks on the extracted data. Rows behind each FAIL/WARN are listed "
        "in 'Check Details'.",
    )
    for check in checks:
        sheet.append(
            [check.code, check.title, check.status, check.summary, check.meaning],
            fills={2: _STATUS_FILLS[check.status]},
        )
    return sheet.rows_written


def _write_check_details(wb: Workbook, checks: list[CheckResult]) -> int:
    sheet = SheetWriter(
        wb,
        "Check Details",
        [
            Col("Code", 8),
            Col("Check", 40, wrap=True),
            Col("Date", 12, DATE),
            Col("Voucher type", 18),
            Col("Voucher no.", 14),
            Col("Ledger / item / party", 36),
            Col("Detail", 60, wrap=True),
            Col("Tally GUID", 40),
            REMARKS,
        ],
    )
    for check in checks:
        for issue in check.issues:
            sheet.append(
                [
                    check.code,
                    check.title,
                    issue.date,
                    issue.voucher_type,
                    issue.voucher_number,
                    issue.reference,
                    issue.detail,
                    issue.guid,
                ]
            )
    if sheet.rows_written == 0:
        sheet.append(["-", "No issues found"])
    return sheet.rows_written


def _write_voucher_types(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Voucher Types",
        [
            Col("Voucher type", 28),
            Col("Vouchers (excl. optional)", 14, INT),
            Col("Optional", 12, INT),
            Col("Cancelled", 12, INT),
            Col("Total Debit (posting)", 20, MONEY),
            Col("Tally count", 14, INT, expert=True),
            Col("Match", 12),
            REMARKS,
        ],
        note=f"Compare with {TALLY_PATHS['statistics']} — period {period_label(data)}. "
        "Tally's Statistics does not count optional vouchers, so they are shown separately. "
        "It also adds sub-types into their built-in type's row (e.g. Journal includes "
        "IC JV Expenses): compare such rows with the sum of the types under them.",
    )
    for summary in voucher_type_summaries(data):
        row = sheet.next_row
        sheet.append(
            [
                summary.voucher_type,
                summary.count - summary.optional,
                summary.optional,
                summary.cancelled,
                summary.debit,
                None,
                _match_formula("B", "F", row),
            ]
        )
    return sheet.rows_written


def _write_day_book(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Day Book",
        [
            Col("Date", 12, DATE),
            Col("Voucher type", 18),
            Col("Voucher no.", 14),
            Col("Party ledger", 32),
            Col("Debit", 16, MONEY),
            Col("Credit", 16, MONEY),
            Col("Narration", 48, wrap=True),
            Col("Reference", 18),
            Col("Cancelled", 10),
            Col("Optional", 10),
            Col("Posts to books", 10),
            Col("Ledger lines", 8, INT),
            Col("Inventory lines", 8, INT),
            Col("Tally GUID", 40),
            REMARKS,
        ],
        note=f"Compare with {TALLY_PATHS['day_book']} — period {period_label(data)}",
    )
    totals = voucher_totals(data)
    for v in data.vouchers:
        debit, credit = totals.get(v.tally_guid, (ZERO, ZERO))
        sheet.append(
            [
                v.date,
                v.voucher_type,
                v.voucher_number,
                v.party_ledger_name,
                debit or None,
                credit or None,
                v.narration,
                v.reference,
                v.is_cancelled,
                v.is_optional,
                counts_in_books(v),
                v.ledger_entry_count,
                v.inventory_entry_count,
                v.tally_guid,
            ]
        )
    return sheet.rows_written


def _bill_allocations_text(allocations: Any) -> str | None:
    if not allocations:
        return None
    parts = []
    for alloc in allocations:
        amount = alloc.get("amount")
        dr, cr = split_dr_cr(Decimal(str(amount)) if amount not in (None, "") else None)
        side = f"{dr:,.2f} Dr" if dr is not None else (f"{cr:,.2f} Cr" if cr is not None else "")
        parts.append(
            " ".join(p for p in (alloc.get("bill_type") or "", alloc.get("name") or "", side) if p)
        )
    return "; ".join(parts)


def _write_voucher_lines(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Voucher Lines",
        [
            Col("Date", 12, DATE),
            Col("Voucher type", 18),
            Col("Voucher no.", 14),
            Col("Line", 6, INT),
            Col("Ledger", 36),
            Col("Debit", 16, MONEY),
            Col("Credit", 16, MONEY),
            Col("Source", 14),
            Col("Stock item", 28),
            Col("Party line", 8),
            Col("Bill allocations", 48, wrap=True),
            Col("Voucher GUID", 40),
            REMARKS,
        ],
        note="Ledger postings of each voucher. 'Item invoice' rows are the Sales/Purchase "
        "ledger amounts Tally keeps on the stock lines of item invoices. Open the voucher in "
        "Tally (Day Book > Enter) to compare.",
    )
    postings = postings_by_voucher(data)
    for v in data.vouchers:
        for p in postings.get(v.tally_guid, []):
            sheet.append(
                [
                    v.date,
                    v.voucher_type,
                    v.voucher_number,
                    p.line_no,
                    p.ledger_name,
                    p.debit,
                    p.credit,
                    p.source,
                    p.stock_item,
                    p.is_party_ledger,
                    _bill_allocations_text(p.bill_allocations),
                    v.tally_guid,
                ]
            )
    return sheet.rows_written


def _write_inventory_lines(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Inventory Lines",
        [
            Col("Date", 12, DATE),
            Col("Voucher type", 18),
            Col("Voucher no.", 14),
            Col("Line", 6, INT),
            Col("Stock item", 36),
            Col("In / Out", 8),
            Col("Billed qty", 12, QTY),
            Col("Actual qty", 12, QTY),
            Col("Unit", 8),
            Col("Rate", 14, MONEY),
            Col("Discount %", 10, QTY),
            Col("Amount", 16, MONEY),
            Col("Godown", 20),
            Col("Batch", 16),
            Col("Accounting ledger", 28),
            Col("Voucher GUID", 40),
            REMARKS,
        ],
        note="Stock item lines of each voucher. 'In' = stock inward (Tally deemed positive).",
    )
    for e in data.inventory_entries:
        direction = (
            None if e.is_deemed_positive is None else ("In" if e.is_deemed_positive else "Out")
        )
        sheet.append(
            [
                e.voucher_date,
                e.voucher_type,
                e.voucher_number,
                e.line_no,
                e.stock_item_name,
                direction,
                _abs(e.billed_quantity if e.billed_quantity is not None else e.quantity),
                _abs(e.actual_quantity),
                e.unit,
                _abs(e.rate),
                e.discount,
                _abs(e.amount),
                e.godown_name,
                e.batch_name,
                e.accounting_ledger,
                e.voucher_guid,
            ]
        )
    return sheet.rows_written


def _abs(value: Decimal | None) -> Decimal | None:
    return None if value is None else abs(value)


def _write_ledgers(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Ledgers",
        [
            Col("Ledger", 36),
            Col("Under (group)", 28),
            Col("Opening Debit", 16, MONEY),
            Col("Opening Credit", 16, MONEY),
            Col("Period Debit", 16, MONEY),
            Col("Period Credit", 16, MONEY),
            Col("Bill-wise", 9),
            Col("GSTIN", 18),
            Col("PAN", 12),
            Col("GST registration", 16),
            Col("State", 16),
            Col("Alias", 20),
            Col("Tally GUID", 40),
            REMARKS,
        ],
        note=f"Period Debit/Credit = 'Current Total' in {TALLY_PATHS['ledger_vouchers']} for "
        f"{period_label(data)}. Opening = balance at 'Books beginning from'.",
    )
    movements = ledger_movements(data)
    for ledger in data.ledgers:
        dr, cr = split_dr_cr(ledger.opening_balance)
        movement = movements.get(ledger.name)
        sheet.append(
            [
                ledger.name,
                ledger.parent_group,
                dr,
                cr,
                movement.debit if movement and movement.debit else None,
                movement.credit if movement and movement.credit else None,
                ledger.is_billwise_on,
                ledger.gstin,
                ledger.pan,
                ledger.gst_registration_type,
                ledger.state,
                ledger.alias,
                ledger.tally_guid,
            ]
        )
    return sheet.rows_written


def _write_groups(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Groups",
        [
            Col("Group", 36),
            Col("Under", 28),
            Col("Revenue (P&L)", 12),
            Col("Affects gross profit", 14),
            Col("Debit nature", 12),
            Col("Sub-ledger", 10),
            Col("Alias", 20),
            Col("Tally GUID", 40),
            REMARKS,
        ],
        note=f"Compare with {TALLY_PATHS['groups']}.",
    )
    for g in data.groups:
        sheet.append(
            [
                g.name,
                g.parent,
                g.is_revenue,
                g.affects_gross_profit,
                g.is_deemed_positive,
                g.is_subledger,
                g.alias,
                g.tally_guid,
            ]
        )
    return sheet.rows_written


def _write_stock_items(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Stock Items",
        [
            Col("Stock item", 36),
            Col("Under (stock group)", 24),
            Col("Category", 18),
            Col("Unit", 8),
            Col("Opening qty", 12, QTY),
            Col("Opening rate", 14, MONEY),
            Col("Opening value", 16, MONEY),
            Col("HSN/SAC", 12),
            Col("Part no.", 14),
            Col("GST applicable", 14),
            Col("Alias", 20),
            Col("Tally GUID", 40),
            REMARKS,
        ],
        note=f"Compare with {TALLY_PATHS['stock_items']}.",
    )
    for s in data.stock_items:
        sheet.append(
            [
                s.name,
                s.parent,
                s.category,
                s.base_unit,
                s.opening_quantity,
                _abs(s.opening_rate),
                _abs(s.opening_value),
                s.hsn_code,
                s.part_number,
                s.gst_applicable,
                s.alias,
                s.tally_guid,
            ]
        )
    return sheet.rows_written


def _write_bills(wb: Workbook, data: CompanyReportData) -> int:
    sheet = SheetWriter(
        wb,
        "Bills Outstanding",
        [
            Col("Receivable / Payable", 14),
            Col("Party ledger", 36),
            Col("Bill ref.", 18),
            Col("Bill date", 12, DATE),
            Col("Due date", 12, DATE),
            Col("Overdue days", 10, INT),
            Col("Pending Debit", 16, MONEY),
            Col("Pending Credit", 16, MONEY),
            Col("Opening amount", 16, MONEY),
            Col("Advance", 9),
            REMARKS,
        ],
        note=f"Compare with {TALLY_PATHS['receivables']} / Payables.",
    )
    for b in data.bills:
        dr, cr = split_dr_cr(b.closing_amount)
        opening_dr, opening_cr = split_dr_cr(b.opening_amount)
        sheet.append(
            [
                b.bill_type.title(),
                b.ledger_name,
                b.bill_ref,
                b.bill_date,
                b.due_date,
                b.overdue_days,
                dr,
                cr,
                opening_dr if opening_dr is not None else opening_cr,
                b.is_advance,
            ]
        )
    return sheet.rows_written


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def write_company_workbook(
    data: CompanyReportData,
    output_dir: Path,
    generated_at: datetime,
) -> ExcelReportResult:
    """Write the verification workbook for one company and return where it went."""
    checks = run_checks(data)
    wb = Workbook(write_only=True)
    wb.properties.title = f"Tally data verification — {data.company.company_name}"
    wb.properties.creator = f"stallion-tally {__version__}"

    result = ExcelReportResult(
        path=Path(),
        company_name=data.company.company_name,
        period_from=data.period_from,
        period_to=data.period_to,
        checks=checks,
    )
    result.sheet_rows["Read Me"] = _write_read_me(wb, data, generated_at)
    result.sheet_rows["Summary"] = _write_summary(wb, data)
    result.sheet_rows["Checks"] = _write_checks(wb, checks)
    result.sheet_rows["Check Details"] = _write_check_details(wb, checks)
    result.sheet_rows["Voucher Types"] = _write_voucher_types(wb, data)
    for name, writer in (
        ("Day Book", _write_day_book),
        ("Voucher Lines", _write_voucher_lines),
        ("Inventory Lines", _write_inventory_lines),
        ("Ledgers", _write_ledgers),
        ("Groups", _write_groups),
        ("Stock Items", _write_stock_items),
        ("Bills Outstanding", _write_bills),
    ):
        result.sheet_rows[name] = writer(wb, data)

    output_dir.mkdir(parents=True, exist_ok=True)
    period = (
        f"_{data.period_from:%Y%m%d}-{data.period_to:%Y%m%d}"
        if data.period_from and data.period_to
        else ""
    )
    name = (
        f"Tally_Verification_{safe_file_part(data.company.company_name)}{period}_"
        f"{_local(generated_at):%Y%m%d-%H%M%S}.xlsx"
    )
    path = output_dir / name
    partial = path.with_suffix(".xlsx.partial")
    wb.save(partial)
    os.replace(partial, path)
    result.path = path
    return result
