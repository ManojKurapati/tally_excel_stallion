from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from stallion_tally.cli.main import cli
from stallion_tally.config import Settings
from stallion_tally.database import Database
from stallion_tally.database.models import CompanyRow, LedgerRow, VoucherRow
from stallion_tally.reports import load_company_report, write_company_workbook
from stallion_tally.reports.checks import split_dr_cr
from stallion_tally.reports.excel import Col, SheetWriter
from stallion_tally.sync.manager import SyncManager
from tests.conftest import COMPANY_1, COMPANY_1_ID, NOW, FakeTally

SALES_GUID = f"{COMPANY_1_ID}-00000101"

SHEETS = [
    "Read Me",
    "Summary",
    "Checks",
    "Check Details",
    "Voucher Types",
    "Day Book",
    "Voucher Lines",
    "Inventory Lines",
    "Ledgers",
    "Groups",
    "Stock Items",
    "Bills Outstanding",
]


def build(database: Database, out: Path, **kwargs):
    with database.session() as session:
        company = session.execute(
            select(CompanyRow).where(CompanyRow.company_name == COMPANY_1)
        ).scalar_one()
        data = load_company_report(session, company, **kwargs)
        return data, write_company_workbook(data, out, NOW)


def summary_values(wb) -> dict[str, object]:
    ws = wb["Summary"]
    return {row[1]: row[2] for row in ws.iter_rows(min_row=3, values_only=True)}


def check_status(wb) -> dict[str, str]:
    return {row[0]: row[2] for row in wb["Checks"].iter_rows(min_row=3, values_only=True)}


def test_split_dr_cr_uses_tally_sign() -> None:
    assert split_dr_cr(Decimal("-100")) == (Decimal("100"), None)
    assert split_dr_cr(Decimal("250.5")) == (None, Decimal("250.5"))
    assert split_dr_cr(Decimal(0)) == (None, None)
    assert split_dr_cr(None) == (None, None)


def test_workbook_matches_local_database(
    manager: SyncManager, database: Database, tmp_path: Path
) -> None:
    assert manager.run(export=False).status == "completed_no_export"
    data, result = build(database, tmp_path / "xl")

    assert result.path.exists() and result.path.suffix == ".xlsx"
    assert COMPANY_1 in result.path.name
    assert not list((tmp_path / "xl").glob("*.partial"))
    assert data.period_source == "last Day Book window synced from Tally"

    wb = load_workbook(result.path)
    assert wb.sheetnames == SHEETS

    with database.session() as s:
        ledgers = s.scalars(
            select(LedgerRow).where(LedgerRow.company_name == COMPANY_1, ~LedgerRow.is_deleted)
        ).all()
        vouchers = s.scalars(
            select(VoucherRow).where(
                VoucherRow.company_name == COMPANY_1,
                ~VoucherRow.is_deleted,
                VoucherRow.date >= data.period_from,
                VoucherRow.date <= data.period_to,
            )
        ).all()
    summary = summary_values(wb)
    assert summary["Number of ledgers"] == len(ledgers)
    assert (
        summary["Number of vouchers (excluding optional)"]
        == sum(not v.is_optional for v in vouchers)
        > 0
    )
    assert result.sheet_rows["Day Book"] == len(vouchers)
    assert summary["Total Debit"] == pytest.approx(summary["Total Credit"])

    # Match formula compares the extracted figure with the expert's entry.
    ws = wb["Summary"]
    assert ws["E3"].value == '=IF(D3="","",IF(ABS(C3-D3)<0.01,"OK","MISMATCH"))'

    guids = {row[13] for row in wb["Day Book"].iter_rows(min_row=3, values_only=True)}
    assert guids == {v.tally_guid for v in vouchers}

    # Item invoice: the Sales ledger credit comes from the stock line's accounting allocation.
    item_rows = [
        row
        for row in wb["Voucher Lines"].iter_rows(min_row=3, values_only=True)
        if row[11] == SALES_GUID and row[7] == "Item invoice"
    ]
    assert item_rows and item_rows[0][6] == pytest.approx(10000)

    status = check_status(wb)
    assert status["C01"] == "PASS"
    assert status["C02"] == "PASS"
    # The fixture ledger master is deliberately partial, so only C04 (unknown ledgers) fails.
    assert {c.code for c in result.checks if c.status == "FAIL"} == {"C04"}


def test_unbalanced_voucher_is_flagged(
    manager: SyncManager, database: Database, fake_tally: FakeTally, tmp_path: Path
) -> None:
    fake_tally.drop_ledger_entry(SALES_GUID)
    manager.run(export=False)
    _, result = build(database, tmp_path)

    wb = load_workbook(result.path)
    assert check_status(wb)["C01"] == "FAIL"
    detail_guids = {row[7] for row in wb["Check Details"].iter_rows(min_row=2, values_only=True)}
    assert SALES_GUID in detail_guids
    assert result.failed_checks >= 1


def test_tally_text_is_never_written_as_formula(
    manager: SyncManager, database: Database, fake_tally: FakeTally, tmp_path: Path
) -> None:
    evil = '=HYPERLINK("http://example.invalid","x")'
    fake_tally.set_narration(SALES_GUID, evil)
    manager.run(export=False)
    _, result = build(database, tmp_path)

    ws = load_workbook(result.path)["Day Book"]
    cells = [row[6] for row in ws.iter_rows(min_row=3) if row[13].value == SALES_GUID]
    assert cells[0].value == evil
    assert cells[0].data_type == "s"


def test_explicit_period_filters_vouchers(
    manager: SyncManager, database: Database, tmp_path: Path
) -> None:
    manager.run(export=False)
    data, _ = build(database, tmp_path, from_date=date(2030, 1, 1), to_date=date(2030, 1, 31))
    assert data.period_source == "requested"
    assert data.vouchers == [] and data.ledger_entries == []


def test_sheet_writer_continues_on_new_sheet(tmp_path: Path) -> None:
    wb = Workbook(write_only=True)
    sheet = SheetWriter(wb, "Day Book", [Col("A"), Col("B")], max_rows=4)
    sheet.extend([i, f"row {i}"] for i in range(5))
    path = tmp_path / "split.xlsx"
    wb.save(path)

    loaded = load_workbook(path)
    # 4 rows per sheet including the header: 3 + 2 data rows
    assert loaded.sheetnames == ["Day Book", "Day Book (2)"]
    values = [
        row[0]
        for name in loaded.sheetnames
        for row in loaded[name].iter_rows(min_row=2, values_only=True)
    ]
    assert values == [0, 1, 2, 3, 4]
    assert loaded["Day Book (2)"]["A1"].value == "A"


def test_cli_excel_command(manager: SyncManager, settings: Settings, tmp_path: Path) -> None:
    manager.run(export=False)
    env_file = tmp_path / "empty.env"
    env_file.write_text("", encoding="utf-8")
    out = tmp_path / "cli_out"

    runner = CliRunner()
    result = runner.invoke(
        cli,
        [
            "--env-file",
            str(env_file),
            "--data-dir",
            str(settings.app_data_dir),
            "excel",
            "--company",
            COMPANY_1,
            "--output",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    files = list(out.glob("*.xlsx"))
    assert len(files) == 1 and COMPANY_1 in files[0].name

    missing = runner.invoke(
        cli,
        [
            "--env-file",
            str(env_file),
            "--data-dir",
            str(settings.app_data_dir),
            "excel",
            "--company",
            "No Such Company",
            "--output",
            str(out),
        ],
    )
    assert missing.exit_code == 1


def _report(vouchers, ledger_entries, inventory_entries):
    from types import SimpleNamespace

    from stallion_tally.reports.loader import CompanyReportData

    data = CompanyReportData(SimpleNamespace(), None, None, "test")
    data.vouchers = [SimpleNamespace(**v) for v in vouchers]
    data.ledger_entries = [SimpleNamespace(**e) for e in ledger_entries]
    data.inventory_entries = [SimpleNamespace(**e) for e in inventory_entries]
    return data


def _vch(guid: str, *, view: str, invoice: bool, raw_keys: tuple[str, ...]) -> dict:
    return dict(
        tally_guid=guid,
        voucher_type="T",
        voucher_number=guid,
        date=date(2026, 1, 1),
        party_ledger_name=None,
        persisted_view=view,
        is_invoice=invoice,
        is_cancelled=False,
        is_optional=False,
        raw_json=dict.fromkeys(raw_keys, []),
        ledger_entry_count=1,
        inventory_entry_count=1,
    )


def _line(guid: str, n: int, ledger: str, amount: str) -> dict:
    debit, credit = split_dr_cr(Decimal(amount))
    return dict(
        voucher_guid=guid,
        line_no=n,
        ledger_name=ledger,
        debit=debit,
        credit=credit,
        bill_allocations_json=None,
        is_party_ledger=None,
    )


def _stock(guid: str, amount: str) -> dict:
    return dict(
        voucher_guid=guid, line_no=1, stock_item_name="Item", accounting_ledger="Sales",
        amount=Decimal(amount),
    )  # fmt: skip


def test_item_invoice_postings_follow_tally_ledger_list() -> None:
    from stallion_tally.reports.checks import (
        check_vouchers_balanced,
        ledger_movements,
        period_totals,
    )

    invoice_view = "Invoice Voucher View"
    data = _report(
        [
            # Collection export: ALLLEDGERENTRIES already holds the Sales ledger.
            _vch("full", view=invoice_view, invoice=True, raw_keys=("ALLLEDGERENTRIES.LIST",)),
            # Older Day Book export: Sales ledger only on the stock line.
            _vch("old", view=invoice_view, invoice=True, raw_keys=("LEDGERENTRIES.LIST",)),
            # Sales order: not an invoice, never posts.
            _vch("order", view=invoice_view, invoice=False, raw_keys=("ALLLEDGERENTRIES.LIST",)),
        ],
        [
            _line("full", 1, "Party", "-100"),
            _line("full", 2, "Sales", "100"),
            _line("old", 1, "Party", "-50"),
            _line("order", 1, "Party", "-70"),
        ],
        [_stock("full", "100"), _stock("old", "50"), _stock("order", "70")],
    )
    assert check_vouchers_balanced(data).status == "PASS"
    assert period_totals(data) == (Decimal("150"), Decimal("150"))
    movements = ledger_movements(data)
    assert movements["Sales"].credit == Decimal("150")
    assert movements["Party"].debit == Decimal("150")
